#!/usr/bin/env python3
"""Build a promotional-badge/overlay classification dataset from product sheets.

For each product's first image we compute:
  - CLIP image features (512-d) + DINOv2 features (384-d)
  - the CLIP "added text overlay" zero-shot score
  - OCR ops-claim text (if any)

Weak label (starting point — replace with human labels for best results):
  1  = overlay present  (OCR claim found OR added_text >= BADGE_POS)
  0  = clean            (added_text <= BADGE_NEG)
  (skipped otherwise)

Usage: python3 _badge_dataset.py <sheet.xlsx> <out.npz> [--limit N]
"""
import os
import re
import sys
import json

import numpy as np

os.environ.setdefault("EMBED_BACKEND", "llama")


def _col(hdr, name):
    for i, h in enumerate(hdr):
        if h.replace("*", "").strip().lower() == name.lower():
            return i
    return -1


def main():
    sheet = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "badge_dataset.npz"
    limit = None
    if "--limit" in sys.argv:
        limit = int(sys.argv[sys.argv.index("--limit") + 1])

    import openpyxl
    from _clip_verify import load_image, get_clip, _get_dinov2, clip_zero_shot, check_ops_on_image, RULE5_CHECKS

    wb = openpyxl.load_workbook(sheet, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c) if c is not None else "" for c in rows[1]]
    si = _col(hdr, "Sku")
    imgi = [_col(hdr, f"Product Image {i}") for i in range(1, 11)]
    urls = []
    for r in rows[2:]:
        if not r or r[si] in (None, ""):
            continue
        for i in imgi:
            if i >= 0 and i < len(r) and r[i] and str(r[i]).strip():
                urls.append((str(r[si]).strip(), str(r[i]).strip()))
                break  # first image only
        if limit and len(urls) >= limit:
            break
    wb.close()
    print(f"[DATA] {len(urls)} first-images to process")

    # added_text prompts
    pos = neg = None
    for name, a, b, t in RULE5_CHECKS:
        if name == "added_text":
            pos, neg = a, b

    clip_model, clip_proc = get_clip()
    dino_model, dino_proc, dino_dev = _get_dinov2()
    import torch

    feats, labels, kept_urls, scores = [], [], [], []
    for n, (sku, u) in enumerate(urls):
        img = load_image(u)
        if img is None:
            continue
        try:
            # CLIP image feature
            ci = {k: (v.to(next(clip_model.parameters()).device, next(clip_model.parameters()).dtype) if v.is_floating_point() else v.to(next(clip_model.parameters()).device))
                  for k, v in clip_proc(images=img, return_tensors="pt").items()}
            with torch.no_grad():
                cf = clip_model.get_image_features(**ci)
            cf = (cf / cf.norm(dim=-1, keepdim=True)).float().cpu().numpy()[0]
            # DINOv2 feature
            di = {k: v.to(dino_dev) for k, v in dino_proc(images=img, return_tensors="pt").items()}
            with torch.no_grad():
                out_d = dino_model(**di)
                df = out_d.last_hidden_state[:, 0, :]
            df = (df / df.norm(dim=-1, keepdim=True)).float().cpu().numpy()[0]
        except Exception as e:
            print(f"[DATA] feature fail {u}: {e}")
            continue
        # scores
        s = clip_zero_shot(img, [pos, neg])[pos]
        r7 = check_ops_on_image(u)
        claim = bool(r7.get("flagged"))
        if claim or s >= 0.85:
            lab = 1
        elif s <= 0.50:
            lab = 0
        else:
            lab = -1  # ambiguous — skip
        feats.append(np.concatenate([cf, df]))
        labels.append(lab)
        kept_urls.append(u)
        scores.append(float(s))
        if (n + 1) % 25 == 0:
            print(f"[DATA] {n+1}/{len(urls)} done")

    feats = np.asarray(feats, dtype=np.float32)
    labels = np.asarray(labels, dtype=np.int64)
    scores = np.asarray(scores, dtype=np.float32)
    mask = labels >= 0
    np.savez(out, feats=feats[mask], labels=labels[mask], scores=scores[mask],
             urls=np.asarray(kept_urls, dtype=object)[mask])
    print(f"[DATA] saved {out}: {int(mask.sum())} labelled "
          f"({int((labels[mask]==1).sum())} overlay / {int((labels[mask]==0).sum())} clean) "
          f"of {len(labels)} total")


if __name__ == "__main__":
    main()
