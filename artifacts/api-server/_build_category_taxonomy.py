#!/usr/bin/env python3
"""
Build the L1->L4 category taxonomy JSON from the "All Categories" workbook.

This is the authoritative single-sheet format:
    Category ID | L1 | L2 | L3 | L4 | Full Path | Level | Sort Order | Status | ...

The workbook already carries the resolved `Full Path` per row, so this builder
is simpler and more reliable than the multi-sheet L1/L2/L3/L4 workbooks: it
reads every row, keeps the served hierarchy, and records status + ids.

Usage:
    python3 _build_category_taxonomy.py "All Categories.xlsx" category_l1l4.json

Output shape (consumed by _category_validate.py):
    {
      "version", "source",
      "l1": [names...],
      "l2": {"<l1>": [names...]},
      "l3": {"<l1> || <l2>": [names...]},
      "l4": {"<l1> || <l2> || <l3>": [names...]},
      "ids": {"<full path>": categoryId},
      "status": {"<full path>": "Active"|"Inactive"},
      "paths": [{l1,l2,l3,l4,full,id,status}, ...]   # ONLY L4 leaves
    }
"""

import json
import sys
from collections import OrderedDict

try:
    import openpyxl
except Exception as e:  # pragma: no cover
    print(f"[TAXONOMY] openpyxl required: {e}", file=sys.stderr)
    sys.exit(1)


def _norm(v) -> str:
    if v is None:
        return ""
    return " ".join(str(v).split()).strip()


def _find_col(header, aliases):
    lower = [str(h).strip().lower() if h is not None else "" for h in header]
    for a in aliases:
        if a in lower:
            return lower.index(a)
    return -1


def build(xlsx_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Categories"] if "Categories" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        wb.close()
        return {"error": "empty workbook"}

    c_id = _find_col(header, ["category id", "id"])
    c_l1 = _find_col(header, ["l1"])
    c_l2 = _find_col(header, ["l2"])
    c_l3 = _find_col(header, ["l3"])
    c_l4 = _find_col(header, ["l4"])
    c_full = _find_col(header, ["full path", "path"])
    c_status = _find_col(header, ["status"])

    def get(row, i):
        return _norm(row[i]) if 0 <= i < len(row) else ""

    l1_col = OrderedDict()
    l2_map = OrderedDict()
    l3_map = OrderedDict()
    l4_map = OrderedDict()
    ids = {}
    status = {}
    paths = []
    seen = set()

    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        a, b, c, d = get(row, c_l1), get(row, c_l2), get(row, c_l3), get(row, c_l4)
        if not a:
            continue

        l1_col.setdefault(a, None)
        if b:
            l2_map.setdefault(a, OrderedDict()).setdefault(b, None)
            k2 = f"{a} || {b}"
            if c:
                l3_map.setdefault(k2, OrderedDict()).setdefault(c, None)
                k3 = f"{a} || {b} || {c}"
                if d:
                    l4_map.setdefault(k3, OrderedDict()).setdefault(d, None)

        full = get(row, c_full) or " > ".join([x for x in (a, b, c, d) if x])
        cid = get(row, c_id)
        st = get(row, c_status)
        if full:
            ids[full] = cid
            status[full] = st

        # Leaves only (a full L4 path) drive classification.
        if a and b and c and d:
            sig = (a, b, c, d)
            if sig not in seen:
                seen.add(sig)
                paths.append({
                    "l1": a, "l2": b, "l3": c, "l4": d,
                    "full": full,
                    "id": cid,
                    "status": st,
                })

    wb.close()
    return {
        "version": "2.0",
        "source": xlsx_path.split("/")[-1],
        "l1": list(l1_col.keys()),
        "l2": {k: list(v.keys()) for k, v in l2_map.items()},
        "l3": {k: list(v.keys()) for k, v in l3_map.items()},
        "l4": {k: list(v.keys()) for k, v in l4_map.items()},
        "ids": ids,
        "status": status,
        "paths": paths,
    }


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python3 _build_category_taxonomy.py <xlsx> [out.json]", file=sys.stderr)
        sys.exit(1)
    xlsx_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else "category_l1l4.json"
    data = build(xlsx_path)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(
        f"[TAXONOMY] {len(data['l1'])} L1, "
        f"{sum(len(v) for v in data['l2'].values())} L2, "
        f"{sum(len(v) for v in data['l3'].values())} L3, "
        f"{len(data['paths'])} L4 paths -> {out_path}",
        file=sys.stderr,
    )


if __name__ == "__main__":
    main()
