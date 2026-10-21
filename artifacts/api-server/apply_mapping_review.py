#!/usr/bin/env python3
"""Fold human corrections from an exported "Mapping Review" sheet back into
marketplace_category_map.json.

Usage:
    python3 apply_mapping_review.py "AHUJA MART.xlsx"
    python3 apply_mapping_review.py "AHUJA MART.xlsx" --dry-run
    python3 apply_mapping_review.py "sheet.xlsx" --depth 2
    python3 apply_mapping_review.py "sheet.xlsx" --exact

How it works:
    1. every exported workbook carries a "Mapping Review" sheet listing the rows
       the mapper had to GUESS (coarse / coarse-cross / above-L4 / unmapped)
    2. type the correct Gajab path into its "Corrected Category" column, e.g.
         Grocery > Food & Nutrition > Spices, Condiments & Sauces > Spices & Masalas
    3. run this script

Corrections are keyed on the SOURCE BREADCRUMB PREFIX (default: the first three
crumbs), so confirming one row fixes the whole family:
    "Grocery & Gourmet Foods > Dried Fruits, Nuts & Seeds > Dried Fruits"
covers Dates, Raisins, Anjeer, Cranberries and every future scrape with the same
breadcrumb. Use --exact to key the full breadcrumb instead (a single row only),
or --depth 2/4 to widen/narrow the prefix.

Targets are validated against category_l1l4.json; a path may be any level that
owns L4 leaves (L1/L2/L3 ancestors and the L4 itself). Nothing is written when
the map would be left in an invalid state.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _category_map import (  # noqa: E402
    _default_map,
    _leafful_sets,
    _norm_key,
    load_taxonomy,
    validate_map,
)

REVIEW_SHEET = "Mapping Review"
COL_PLATFORM = "Platform"
COL_SOURCE = "Source Category"
COL_CORRECTED = "Corrected Category"


def _header_map(header_row: tuple) -> dict[str, int]:
    """Column index by collapsed header name (template headers carry double spaces)."""
    return {" ".join(str(h).split()): i for i, h in enumerate(header_row) if h}


def _cell(row: tuple, idx: dict[str, int], name: str) -> str:
    i = idx.get(name)
    if i is None or i >= len(row) or row[i] is None:
        return ""
    return str(row[i]).strip()


def _prefix(source: str, depth: int, exact: bool) -> str:
    parts = [p.strip() for p in source.split(" > ") if p.strip()]
    return source.strip() if exact else " > ".join(parts[: max(1, depth)])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("workbook", help="exported .xlsx containing a 'Mapping Review' sheet")
    ap.add_argument("--map", default=None, help="marketplace_category_map.json (default: next to this script)")
    ap.add_argument("--platform", default="", help="fallback platform when the column is blank")
    ap.add_argument("--depth", type=int, default=3, help="breadcrumb prefix depth to key on (default 3)")
    ap.add_argument("--exact", action="store_true", help="key the FULL breadcrumb instead of a prefix")
    ap.add_argument("--dry-run", action="store_true", help="report what would change, write nothing")
    args = ap.parse_args()

    wb_path = Path(args.workbook).expanduser()
    if not wb_path.exists():
        print(f"error: {wb_path} not found", file=sys.stderr)
        return 2
    wb = load_workbook(wb_path, read_only=True, data_only=True)
    if REVIEW_SHEET not in wb.sheetnames:
        print(f"error: no '{REVIEW_SHEET}' sheet - re-export with the current build", file=sys.stderr)
        print(f"       sheets: {', '.join(wb.sheetnames[:8])} ...", file=sys.stderr)
        return 2

    rows = list(wb[REVIEW_SHEET].iter_rows(values_only=True))
    if len(rows) < 2:
        print("no review rows")
        return 0
    idx = _header_map(rows[0])

    taxonomy = load_taxonomy()
    _, valid, _ = _leafful_sets(taxonomy)

    map_path = Path(args.map).expanduser() if args.map else Path(_default_map())
    data = json.loads(map_path.read_text(encoding="utf-8"))
    bucket = data.setdefault("map", {}).setdefault("amazon", {})
    buckets = data["map"]

    applied, invalid, skipped, no_prefix = [], [], [], []
    for r in rows[1:]:
        corrected = _cell(r, idx, COL_CORRECTED)
        if not corrected:
            continue
        platform = (_cell(r, idx, COL_PLATFORM) or args.platform).lower().strip()
        source = _cell(r, idx, COL_SOURCE)
        if corrected not in valid:
            invalid.append((corrected, source[:80]))
            continue
        if not platform:
            skipped.append((corrected, "no platform column value (pass --platform)"))
            continue
        if not source:
            no_prefix.append((corrected, _cell(r, idx, "Product Name")[:60]))
            continue
        key = _prefix(source, args.depth, args.exact)
        target_bucket = buckets.setdefault(platform, {})
        existed = _norm_key(key) in {_norm_key(k) for k in target_bucket}
        # drop any stale spelling of the same key so validate_map does not keep both
        for k in [k for k in target_bucket if _norm_key(k) == _norm_key(key)]:
            del target_bucket[k]
        target_bucket[key] = corrected
        applied.append((platform, key, corrected, existed))

    # validate_map() normalises every key, which would rewrite the whole file
    # into unreadable form; use it as a CHECK only and write the original.
    _, errors = validate_map(data, taxonomy)
    if errors:
        print("refusing to write: the map would be invalid", file=sys.stderr)
        for e in errors[:10]:
            print("  " + e, file=sys.stderr)
        return 1

    for platform, key, corrected, existed in applied:
        verb = "updated" if existed else "added  "
        print(f"  {verb} [{platform}] {key}\n          -> {corrected}")
    for corrected, source in invalid:
        print(f"  SKIP invalid target {corrected!r} (row: {source})")
    for corrected, why in skipped:
        print(f"  SKIP {corrected!r}: {why}")
    for corrected, name in no_prefix:
        print(f"  SKIP {corrected!r}: no source breadcrumb ({name}) - add a title keyword instead")

    if not applied:
        print("\nnothing to apply")
        return 0

    if args.dry_run:
        print(f"\nDRY RUN - {len(applied)} correction(s) would be written to {map_path}")
        return 0

    map_path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"\nwrote {len(applied)} correction(s) to {map_path}")
    print("restart the API (or re-copy the file into dist/) so the export picks it up")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
