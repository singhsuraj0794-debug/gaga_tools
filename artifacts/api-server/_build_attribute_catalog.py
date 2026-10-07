#!/usr/bin/env python3
"""Build gajab_attributes.json from the master attributes workbook.

Source of truth: "Gajab_Attribute_Groups_Final_v3.xlsx", sheet "L4 x Attribute"
(58k+ rows, one row per (L4 node, attribute)). The legacy sheet "Attribute Final"
of "Attributes Final With Valid Entries.xlsx" is also supported.

Columns are matched by header name (not position), so either workbook works:
  L1, L2, L3, L4, Node_Path, Attribute, AttributeType, Description,
  Attribute Group Name, Attribute Specification, Mandatory,
  Valid Entries (Checkbox/Dropdown)

Output shape (unchanged):
  {
    "version": "1.0",
    "source": "<file>",
    "by_path": { "<L1 > L2 > L3 > L4>": [ {attribute,type,mandatory,
                 specification,leaf,group,description,valid}, ... ] },
    "columns": [...], "attribute_names": {...}, "stats": {...}
  }

Usage:
    python3 _build_attribute_catalog.py "Gajab_Attribute_Groups_Final_v3.xlsx" \
        gajab_attributes.json --sheet "L4 x Attribute"
"""

from __future__ import annotations

import argparse
import json
from collections import OrderedDict
from pathlib import Path

import openpyxl

DEFAULT_SHEET = "Attribute Final"


def _clean(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\u00a0", " ").strip()


def _split_valid(raw: str) -> list[str]:
    """Valid entries are ';'-separated (sometimes with a 'Values Accepted - ' prefix)."""
    if not raw:
        return []
    parts = []
    for chunk in raw.split(";"):
        v = chunk.strip()
        if not v:
            continue
        if v.lower().startswith("values accepted - "):
            v = v.split("-", 1)[1].strip()
        if v and v not in parts:
            parts.append(v)
    return parts


def _header_index(header: list[str]) -> dict[str, int]:
    """Map canonical field -> column index, by fuzzy header match."""
    norm = [(_clean(h)).lower() for h in header]
    idx: dict[str, int] = {}
    wanted = {
        "l1": lambda h: h == "l1",
        "l2": lambda h: h == "l2",
        "l3": lambda h: h == "l3",
        "l4": lambda h: h == "l4",
        "node_path": lambda h: h == "node_path",
        "attribute": lambda h: h == "attribute",
        "type": lambda h: h in ("attributetype", "attribute type"),
        "description": lambda h: h == "description",
        "group": lambda h: h.startswith("attribute group name"),
        "specification": lambda h: h.startswith("attribute specification"),
        "mandatory": lambda h: h == "mandatory",
        "valid": lambda h: h.startswith("valid entries"),
    }
    for field, test in wanted.items():
        for i, h in enumerate(norm):
            if test(h):
                idx[field] = i
                break
    return idx


def build(xlsx_path: str, sheet: str = DEFAULT_SHEET) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"sheet {sheet!r} not found; have {wb.sheetnames}")
    ws = wb[sheet]

    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    ix = _header_index(header)
    for req in ("attribute", "l1"):
        if req not in ix:
            raise SystemExit(f"header missing {req!r}; got {header}")

    def get(row, field, default=""):
        i = ix.get(field)
        return _clean(row[i]) if i is not None and i < len(row) else default

    by_path: "OrderedDict[str, list]" = OrderedDict()
    columns: "OrderedDict[str, dict]" = OrderedDict()
    names: dict[str, dict] = {}
    rows = 0
    skipped = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        l1, l2, l3, l4 = (get(row, k) for k in ("l1", "l2", "l3", "l4"))
        node_path = get(row, "node_path")
        attribute = get(row, "attribute")
        attr_type = get(row, "type")
        description = get(row, "description")
        group_name = get(row, "group")
        specification = get(row, "specification")
        mandatory_raw = get(row, "mandatory")
        valid_raw = get(row, "valid")

        # Build the path from the L1..L4 columns — the authoritative current
        # names. The workbook's Node_Path column can carry a stale L1 (e.g.
        # 'Kids & Baby' for what is now 'Toys & General Merchandise'), which
        # would never match the taxonomy.
        joined = " > ".join(p for p in (l1, l2, l3, l4) if p)
        node_path = joined if joined.count(">") >= 1 else node_path

        if not node_path or not attribute:
            skipped += 1
            continue

        # Attribute Specification is the FAMILY name: the L4 name, or the L3
        # name when L4 is blank (the L3-leaf families, e.g.
        # 'Toys & General Merchandise > Toys > Action Figures'). A few cells are
        # a broken '#N/A' — drop those rows.
        if specification.startswith("#"):
            skipped += 1
            continue
        leaf = specification or l4 or l3
        entry = {
            "attribute": attribute,
            "type": attr_type,
            "mandatory": mandatory_raw.lower().startswith("y"),
            # The sheet's attribute group-row reads "Specifications - <family> |
            # <attribute>"; Attribute Specification holds only the family, so
            # build the composite here (the builder used to pass it through).
            "specification": f"{leaf} | {attribute}",
            "leaf": leaf,
            "group": group_name,
            "description": description,
            "valid": _split_valid(valid_raw),
        }

        bucket = by_path.setdefault(node_path, [])
        if not any(e["attribute"] == attribute for e in bucket):
            bucket.append(entry)

        spec_key = entry["specification"]
        if spec_key not in columns:
            columns[spec_key] = {
                "specification": spec_key,
                "attribute": attribute,
                "leaf": leaf,
                "type": attr_type,
                "valid": entry["valid"],
                "paths": [node_path],
            }
        elif node_path not in columns[spec_key]["paths"]:
            columns[spec_key]["paths"].append(node_path)

        nm = names.setdefault(attribute, {"type": attr_type, "valid": []})
        for v in entry["valid"]:
            if v not in nm["valid"]:
                nm["valid"].append(v)

        rows += 1

    wb.close()

    return {
        "version": "1.0",
        "source": Path(xlsx_path).name,
        "sheet": sheet,
        "by_path": by_path,
        "columns": list(columns.values()),
        "attribute_names": names,
        "stats": {
            "rows": rows,
            "skipped": skipped,
            "paths": len(by_path),
            "spec_columns": len(columns),
            "attribute_names": len(names),
        },
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    args = ap.parse_args()

    out = build(args.xlsx, args.sheet)
    dest = Path(args.out)
    dest.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")

    # Slim, gzipped runtime copy: assignment only needs attribute/type/
    # specification/leaf/mandatory per path plus the deduped valid entries.
    import gzip as _gz
    slim = {
        "version": out["version"],
        "by_path": {
            p: [
                {k: a[k] for k in ("attribute", "type", "specification", "leaf", "mandatory")}
                for a in attrs
            ]
            for p, attrs in out["by_path"].items()
        },
        "attribute_names": out["attribute_names"],
    }
    rt = dest.with_name(dest.stem + ".runtime.json.gz")
    rt.write_bytes(_gz.compress(json.dumps(slim, ensure_ascii=False).encode("utf-8")))
    print(f"wrote {rt} ({rt.stat().st_size/1e6:.2f} MB gz)")
    print(f"wrote {dest}")
    s = out["stats"]
    for k in ("rows", "skipped", "paths", "spec_columns", "attribute_names"):
        print(f"  {k:18}: {s[k]}")


if __name__ == "__main__":
    main()
