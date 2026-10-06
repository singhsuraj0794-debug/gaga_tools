#!/usr/bin/env python3
"""
Bulk-create categories directly via the Gajab backend API.

Reads All_category_sheet.xlsx (tabs: L1, L2, L3, L4), skips categories that
already exist (matched by full path), and POSTs the rest to
/product/api/category using the same payload shape as the admin UI.

Usage:
  python3 bulk_create_api.py            # create everything missing
  python3 bulk_create_api.py --limit 5  # create at most 5 (dry test)
  python3 bulk_create_api.py --dry-run  # show what would be created
"""
from __future__ import annotations
import base64, json, re, sys, time, urllib.error, urllib.request
from pathlib import Path
import openpyxl

SCRIPT_DIR = Path(__file__).parent
TOKEN_FILE = SCRIPT_DIR / ".admin_token"
EXCEL_FILE = Path("/Users/gajabmarketing/Downloads/All_category_sheet.xlsx")
IMAGE_FILE = SCRIPT_DIR / "data" / "category_image.jpg"

API = "https://gatewayservice.gajab.com"
ORIGIN = "https://admin.gajab.com"

RETURN_IDS = {"No Return": 5, "Return & Refund": 4}
SHIPPING_IDS = {"Shipping Policy": 3}
BRAND_IDS = {
    "Trademark Certificate": 2,
    "Brand Approval Letter": 1,
    "FSSAI License": 3,
}
DEFAULT_BRAND = [2, 1]


def log(msg: str):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def token() -> str:
    return TOKEN_FILE.read_text().strip()


def api_get(path: str, tok: str):
    req = urllib.request.Request(API + path)
    req.add_header("Authorization", tok)
    req.add_header("Origin", ORIGIN)
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read())


def api_post(path: str, payload: dict, tok: str):
    data = json.dumps(payload).encode()
    req = urllib.request.Request(API + path, data=data, method="POST")
    req.add_header("Authorization", tok)
    req.add_header("Origin", ORIGIN)
    req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            body = r.read()
            return r.status, (json.loads(body) if body else {})
    except urllib.error.HTTPError as e:
        return e.code, {"error": e.read().decode(errors="replace")[:500]}


def slugify(name: str) -> str:
    s = name.lower().replace("&", " ")
    s = re.sub(r"[^a-z0-9\s-]", "", s)
    s = re.sub(r"\s+", "-", s.strip())
    s = re.sub(r"-+", "-", s)
    return s.strip("-")


def fetch_items(tok: str) -> list:
    """Fetch ALL categories with pagination (the endpoint caps a single page)."""
    out, seen, offset, page = [], set(), 0, 1000
    while True:
        d = api_get(f"/product/api/category?limit={page}&offset={offset}&keyword=&status=&name=", tok)
        chunk = d.get("data", [])
        for it in chunk:
            cid = it.get("categoryId")
            if cid not in seen:
                seen.add(cid)
                out.append(it)
        if len(chunk) < page:
            break
        offset += page
    return out


def g(row, i) -> str:
    return str(row[i]).strip() if len(row) > i and row[i] is not None else ""


def parse_sheet() -> dict:
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    cats: dict[str, dict] = {}

    for r in wb["L1"].iter_rows(min_row=2, values_only=True):
        n = g(r, 0)
        if not n or n == "-":
            continue
        cats[n] = dict(level=1, name=n, parent_path="", description=g(r, 2) or f"{n} - L1",
                       commission=g(r, 3) or "30", bargain=g(r, 4) or "25",
                       shipping="", ret="", brand="", status=g(r, 5) or "Enabled")

    for r in wb["L2"].iter_rows(min_row=2, values_only=True):
        l1, l2 = g(r, 0), g(r, 2)
        if not l2:
            continue
        cats[f"{l1} > {l2}"] = dict(level=2, name=l2, parent_path=l1,
                                    description=f"{l2} - L2", commission=g(r, 4) or "30",
                                    bargain=g(r, 5) or "25", shipping=g(r, 6), ret=g(r, 7),
                                    brand=g(r, 8), status=g(r, 9) or "Enabled")

    for r in wb["L3"].iter_rows(min_row=2, values_only=True):
        l1, l2, l3, parent = g(r, 0), g(r, 2), g(r, 4), g(r, 5)
        if not l3:
            continue
        path = f"{parent} > {l3}" if parent else f"{l1} > {l2} > {l3}"
        cats[path] = dict(level=3, name=l3, parent_path=parent,
                          description=f"{l3} - L3", commission=g(r, 6) or "30",
                          bargain=g(r, 7) or "25", shipping=g(r, 8), ret=g(r, 9),
                          brand=g(r, 10), status=g(r, 11) or "Enabled")

    for r in wb["L4"].iter_rows(min_row=2, values_only=True):
        l1, l2, l3, l4, parent = g(r, 0), g(r, 2), g(r, 4), g(r, 6), g(r, 7)
        if not l4:
            continue
        path = f"{parent} > {l4}" if parent else f"{l1} > {l2} > {l3} > {l4}"
        cats[path] = dict(level=4, name=l4, parent_path=parent,
                          description=f"{l4} - L4", commission=g(r, 8) or "30",
                          bargain=g(r, 9) or "25", shipping=g(r, 10), ret=g(r, 11),
                          brand=g(r, 12), status=g(r, 13) or "Enabled")

    return cats


def brand_ids(brand_str: str) -> list:
    ids = []
    for part in brand_str.split(","):
        part = part.strip()
        if part in BRAND_IDS:
            ids.append(BRAND_IDS[part])
    return ids or DEFAULT_BRAND


def build_payload(cat: dict, parent_id, sort_order: int, image_b64: str) -> dict:
    status = 1 if cat["status"].lower() == "enabled" else 0
    ret_id = RETURN_IDS.get(cat["ret"], 4)
    ship_ids = [SHIPPING_IDS.get(cat["shipping"], 3)]
    return {
        "name": cat["name"],
        "sortOrder": sort_order,
        "metaTagDescription": "",
        "metaTagKeyword": "",
        "metaTagTitle": "",
        "parentInt": parent_id,
        "image": image_b64,
        "imagePath": "",
        "status": status,
        "categorySlug": slugify(cat["name"]),
        "categoryDescription": cat["description"],
        "categoryTax": "",
        "isVariant": 0,
        "categoryComission": float(cat["commission"]),
        "bargainPercentage": float(cat["bargain"]),
        "shippingCharge": 0,
        "subImage": [],
        "returnPolicyId": ret_id,
        "shippingPolicyId": ship_ids,
        "documentIds": brand_ids(cat["brand"]),
        "allowLowStockAlert": 0,
    }


def main():
    limit = None
    dry = "--dry-run" in sys.argv
    if "--limit" in sys.argv:
        limit = int(sys.argv[sys.argv.index("--limit") + 1])

    tok = token()
    log("Fetching existing categories...")
    items = fetch_items(tok)
    path_to_id = {i.get("levels", ""): i.get("categoryId") for i in items}
    # max sort order per parent id (top-level uses parent id None -> key 0)
    children_sort: dict = {}
    for i in items:
        pid = i.get("parentInt") or 0
        children_sort[pid] = max(children_sort.get(pid, 0), i.get("sortOrder") or 0)
    log(f"  {len(items)} existing categories")

    cats = parse_sheet()
    log(f"  {len(cats)} unique paths in sheet")

    todo = sorted([p for p in cats if p not in path_to_id],
                  key=lambda p: (cats[p]["level"], p))
    log(f"  {len(todo)} missing")

    if dry:
        for p in todo[:50]:
            log(f"  would create L{cats[p]['level']}: {p}")
        log(f"  ... total {len(todo)}")
        return

    image_b64 = "data:image/jpeg;base64," + base64.b64encode(IMAGE_FILE.read_bytes()).decode()

    created = failed = 0
    failures = []
    for idx, path in enumerate(todo):
        if limit is not None and created >= limit:
            log(f"Reached limit {limit}.")
            break

        cat = cats[path]
        parent_path = cat["parent_path"]
        if parent_path:
            parent_id = path_to_id.get(parent_path)
            if parent_id is None:
                log(f"SKIP (parent missing): {path}")
                failed += 1
                failures.append((path, "parent missing"))
                continue
        else:
            parent_id = None

        pid_key = parent_id or 0
        sort_order = children_sort.get(pid_key, 0) + 1
        payload = build_payload(cat, parent_id, sort_order, image_b64)

        status, resp = api_post("/product/api/category", payload, tok)
        msg = (resp.get("message") or resp.get("error") or "")
        # Backend enforces unique sort order per parent; bump and retry if taken.
        attempts = 0
        while "sort order" in msg.lower() and attempts < 300:
            attempts += 1
            sort_order += 1
            payload["sortOrder"] = sort_order
            status, resp = api_post("/product/api/category", payload, tok)
            msg = (resp.get("message") or resp.get("error") or "")

        ok = status in (200, 201) and (resp.get("status") in (1, "1", True) or "data" in resp)
        if ok:
            created += 1
            children_sort[pid_key] = sort_order
            # discover new id
            new_id = None
            data = resp.get("data")
            if isinstance(data, dict):
                new_id = data.get("categoryId") or data.get("id")
            elif isinstance(data, list) and data and isinstance(data[0], dict):
                new_id = data[0].get("categoryId") or data[0].get("id")
            if new_id is None:
                # fall back to a lookup
                try:
                    fresh = fetch_items(tok)
                    for i in fresh:
                        if i.get("levels") == path:
                            new_id = i.get("categoryId")
                            break
                except Exception:
                    pass
            if new_id is not None:
                path_to_id[path] = new_id
            log(f"[{created}] L{cat['level']} OK  {path}  (id={new_id}, sort={sort_order})")
        else:
            failed += 1
            msg = resp.get("message") or resp.get("error") or str(resp)[:200]
            failures.append((path, msg))
            log(f"FAIL L{cat['level']} {path}: {msg}")

    log("=" * 60)
    log(f"Created {created}, failed {failed}")
    for p, m in failures[:40]:
        log(f"  FAIL {p}: {m}")


if __name__ == "__main__":
    main()
