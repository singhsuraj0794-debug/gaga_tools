#!/usr/bin/env python3
"""
Create Attribute Groups + Specifications and assign specs to L4 families from a
sheet in the NEW format (columns include 'Attribute Group Name' and
'Attribute Specification'). Attributes must already exist.

Model (confirmed against the live admin):
  * Attribute Group Name   -> GLOBAL group; each attribute belongs to one group.
                              Created with the union of its attributes.
  * Attribute Specification -> named after the L2; contains the groups (with their
                              attributes) and is assigned to the L4 families.

Usage:
  python3 create_groups_specs.py --file "/path/sheet.xlsx" [--workers 8] [--dry-run]
"""
from __future__ import annotations
import json, sys, threading, time, urllib.error, urllib.request
from collections import OrderedDict, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
import openpyxl

SCRIPT_DIR = Path(__file__).parent
TOKEN_FILE = SCRIPT_DIR / ".admin_token"
API = "https://gatewayservice.gajab.com"
_TL = threading.Lock()


def log(m):
    with _TL:
        print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


def token():
    return TOKEN_FILE.read_text().strip()


def api(method, path, payload=None, tok=None, tries=5, timeout=90):
    data = json.dumps(payload).encode() if payload is not None else None
    for i in range(tries):
        r = urllib.request.Request(API + path, data=data, method=method)
        for k, v in [("Authorization", tok), ("Origin", "https://admin.gajab.com"),
                     ("Referer", "https://admin.gajab.com/"),
                     ("Accept", "application/json, text/plain, */*")]:
            r.add_header(k, v)
        if data:
            r.add_header("Content-Type", "application/json")
        try:
            resp = urllib.request.urlopen(r, timeout=timeout)
            return resp.status, json.loads(resp.read() or b"{}")
        except urllib.error.HTTPError as e:
            body = e.read().decode(errors="replace")
            try:
                j = json.loads(body)
            except Exception:
                j = {"error": body[:200]}
            if e.code < 500:
                return e.code, j
            time.sleep(1 + i)
        except Exception:
            time.sleep(1 + i)
    return 0, {"error": "exhausted retries"}


def g(row, i):
    return str(row[i]).strip() if len(row) > i and row[i] is not None else ""


def parse(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # pick the data sheet (prefer 'Sheet1', else the one whose header has our columns)
    ws = None
    for name in (["Sheet1"] + wb.sheetnames):
        if name not in wb.sheetnames:
            continue
        cand = wb[name]
        hdr = next(cand.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if "Attribute Group Name" in [g(hdr, i) for i in range(len(hdr))]:
            ws = cand
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
    header = [g(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)), i) for i in range(11)]
    idx = {h: i for i, h in enumerate(header)}
    # locate columns robustly
    ci = {
        "l1": idx.get("L1", 0), "l2": idx.get("L2", 1), "l3": idx.get("L3", 2),
        "l4": idx.get("L4", 3), "path": idx.get("Node_Path", 4), "attr": idx.get("Attribute", 5),
        "mand": idx.get("Mandatory", 7), "group": idx.get("Attribute Group Name", 9),
        "spec": idx.get("Attribute Specification", 10),
    }
    group_attrs = OrderedDict()                       # group -> [attr names]
    spec_group_attrs = defaultdict(lambda: defaultdict(list))  # spec -> group -> [attrs]
    spec_groups = defaultdict(list)                   # spec -> [groups]
    mand = {}                                          # (spec,group,attr) -> bool
    spec_families = defaultdict(set)                  # spec -> {node_path}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[ci["l1"]]:
            continue
        attr, grp, spec = g(r, ci["attr"]), g(r, ci["group"]), g(r, ci["spec"])
        p = g(r, ci["path"])
        m = g(r, ci["mand"]).lower() == "yes"
        if not attr or not grp:
            continue
        group_attrs.setdefault(grp, [])
        if attr not in group_attrs[grp]:
            group_attrs[grp].append(attr)
        if spec:
            if grp not in spec_groups[spec]:
                spec_groups[spec].append(grp)
            if attr not in spec_group_attrs[spec][grp]:
                spec_group_attrs[spec][grp].append(attr)
            mand[(spec, grp, attr)] = m
        if spec and p:
            spec_families[spec].add(p)
    return group_attrs, spec_groups, spec_group_attrs, mand, spec_families


def main():
    args = sys.argv
    path = None
    workers = 8
    dry = "--dry-run" in args
    if "--file" in args:
        path = args[args.index("--file") + 1]
    if "--workers" in args:
        workers = int(args[args.index("--workers") + 1])
    if not path:
        print("need --file"); sys.exit(1)

    tok = token()
    t0 = time.time()
    group_attrs, spec_groups, spec_group_attrs, mand, spec_families = parse(path)
    log(f"groups={len(group_attrs)} specs={len(spec_groups)} "
        f"families={sum(len(v) for v in spec_families.values())}")

    attrs = api("GET", "/product/api/attribute?limit=1000&offset=0&keyword=&count=0", tok=tok)[1]["data"]
    attr_by_name = {a["name"].strip().lower(): a["id"] for a in attrs}
    missing_attrs = {a for g_ in group_attrs.values() for a in g_ if a.lower() not in attr_by_name}
    if missing_attrs:
        log(f"WARNING attributes missing in admin: {sorted(missing_attrs)}")

    cats = api("GET", "/product/api/category?limit=5000&offset=0&keyword=&status=&name=", tok=tok)[1]["data"]
    cat_by_path = {c.get("levels", "").strip(): c["categoryId"] for c in cats}

    existing_groups = {x["name"].strip(): x for x in
                       api("GET", "/product/api/attribute-group?limit=5000&offset=0&keyword=&count=0", tok=tok)[1].get("data", [])}
    existing_specs = {x["name"].strip(): x for x in
                      api("GET", "/product/api/specification?limit=5000&offset=0&keyword=&count=0", tok=tok)[1].get("data", [])}
    log(f"existing groups={len(existing_groups)} specs={len(existing_specs)}")

    if dry:
        for gname, ga in group_attrs.items():
            log(f"  GROUP {gname!r} -> {ga}")
        for spec, gs in spec_groups.items():
            log(f"  SPEC {spec!r}: groups={gs} families={len(spec_families[spec])}")
        return

    # ---- Phase 1: groups (upsert: update existing to full union) ----
    group_id = {n: o["id"] for n, o in existing_groups.items()}
    log(f"Phase 1: upsert {len(group_attrs)} groups")
    gsort = [max((o.get("sortOrder") or 0) for o in existing_groups.values()), 0]

    def mk_group(gn):
        ids = [attr_by_name[a.lower()] for a in group_attrs[gn] if a.lower() in attr_by_name]
        obj = existing_groups.get(gn)
        if obj:
            st, resp = api("PUT", f"/product/api/attribute-group/{obj['id']}",
                           {"name": gn, "sortOrder": obj.get("sortOrder") or 1, "attributeIds": ids}, tok=tok)
            return gn, st, obj["id"]
        with _TL:
            gsort[0] += 1
            so = gsort[0]
        st, resp = api("POST", "/product/api/attribute-group/",
                       {"name": gn, "sortOrder": so, "attributeIds": ids}, tok=tok)
        gid = (resp.get("data") or {}).get("id") if isinstance(resp.get("data"), dict) else None
        return gn, st, gid

    with ThreadPoolExecutor(max_workers=workers) as ex:
        for gn, st, gid in [f.result() for f in as_completed([ex.submit(mk_group, gn) for gn in group_attrs])]:
            if gid:
                group_id[gn] = gid
                log(f"  group {gn!r} -> {gid} ({st})")
            else:
                log(f"  GROUP FAIL {gn!r}: {st}")

    # ---- Phase 2: specs (upsert) ----
    spec_id = {n: o["id"] for n, o in existing_specs.items()}
    log(f"Phase 2: upsert {len(spec_groups)} specs")

    def mk_spec(spec):
        attributeGroup = []
        for grp in spec_groups[spec]:
            gid = group_id.get(grp)
            if not gid:
                continue
            aids = [{"attributeId": attr_by_name[a.lower()],
                     "isMandatory": 1 if mand.get((spec, grp, a)) else 0}
                    for a in spec_group_attrs[spec][grp] if a.lower() in attr_by_name]
            attributeGroup.append({"attributeGroupId": gid, "attributeIds": aids})
        obj = existing_specs.get(spec)
        payload = {"name": spec, "status": 1, "attributeGroup": attributeGroup, "attributeGroupId": []}
        if obj:
            st, resp = api("PUT", f"/product/api/specification/{obj['id']}", payload, tok=tok)
            return spec, st, obj["id"]
        st, resp = api("POST", "/product/api/specification", payload, tok=tok)
        sid = (resp.get("data") or {}).get("id") if isinstance(resp.get("data"), dict) else None
        return spec, st, sid

    with ThreadPoolExecutor(max_workers=workers) as ex:
        for spec, st, sid in [f.result() for f in as_completed([ex.submit(mk_spec, s) for s in spec_groups])]:
            if sid:
                spec_id[spec] = sid
                log(f"  spec {spec!r} -> {sid} ({st})")
            else:
                log(f"  SPEC FAIL {spec!r}: {st}")

    # ---- Phase 3: assignments ----
    log("Phase 3: assign specs to families")
    tasks = []
    for spec, paths in spec_families.items():
        sid = spec_id.get(spec)
        if not sid:
            continue
        for p in paths:
            cid = cat_by_path.get(p) or cat_by_path.get(p.replace("Books & General Merchandise", "Toys & General Merchandise"))
            if cid:
                tasks.append((cid, sid))

    def do_assign(task):
        cid, sid = task
        st, resp = api("PUT", f"/product/api/spec-category/{cid}",
                       {"specification": [{"specificationId": sid}], "deleteSpecificationIds": []}, tok=tok)
        return cid, sid, st in (200, 201)

    done = ok = 0
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for cid, sid, good in [f.result() for f in as_completed([ex.submit(do_assign, t) for t in tasks])]:
            done += 1
            if good:
                ok += 1
            else:
                log(f"  ASSIGN FAIL cat={cid} spec={sid}")
            if done % 500 == 0:
                log(f"  families {done}/{len(tasks)} (ok {ok})")
    log(f"assignments ok={ok}/{len(tasks)}")
    log(f"DONE in {(time.time()-t0)/60:.1f} min")


if __name__ == "__main__":
    main()
