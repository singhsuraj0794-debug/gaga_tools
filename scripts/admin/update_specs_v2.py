#!/usr/bin/env python3
"""
v2 attribute-specification migration.

Source of truth: Gajab_Attribute_Groups_Final_v2.xlsx
  - Attribute Specification column now holds L4 names (per-family specs).
  - 2,767 unique specs.

Plan (confirmed with user):
  1. Upsert the 36 attribute groups from the sheet (union of attributes).
  2. Update the existing 85 spec records with the FIRST 85 specs (by sheet
     order of appearance), then create the remaining specs.
  3. Re-assign families: set each family's correct new spec and remove any
     stale (old per-L2) specs it currently carries.

Usage:
  python3 update_specs_v2.py --file <xlsx> [--workers 16] [--dry-run]
"""
from __future__ import annotations
import json, sys, threading, time, urllib.error, urllib.request
from collections import OrderedDict, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
import openpyxl

SCRIPT_DIR = Path(__file__).parent
TOKEN_FILE = SCRIPT_DIR / ".admin_token"
API = "https://gatewayservice.gajab.com"
_TL = threading.Lock()


def log(m):
    with _TL:
        print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


def token():
    return TOKEN_FILE.read_text().strip()


def api(method, path, payload=None, tok=None, tries=5, timeout=90):
    data = json.dumps(payload).encode() if payload is not None else None
    for i in range(tries):
        r = urllib.request.Request(API + path, data=data, method=method)
        for k, v in [("Authorization", tok), ("Origin", "https://admin.gajab.com"),
                     ("Referer", "https://admin.gajab.com/"),
                     ("Accept", "application/json, text/plain, */*")]:
            r.add_header(k, v)
        if data:
            r.add_header("Content-Type", "application/json")
        try:
            resp = urllib.request.urlopen(r, timeout=timeout)
            return resp.status, json.loads(resp.read() or b"{}")
        except urllib.error.HTTPError as e:
            body = e.read().decode(errors="replace")
            try:
                j = json.loads(body)
            except Exception:
                j = {"error": body[:200]}
            if e.code < 500:
                return e.code, j
            time.sleep(1 + i)
        except Exception:
            time.sleep(1 + i)
    return 0, {"error": "exhausted retries"}


def g(row, i):
    return str(row[i]).strip() if len(row) > i and row[i] is not None else ""


def parse(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = None
    for name in (["Sheet1"] + wb.sheetnames):
        if name not in wb.sheetnames:
            continue
        cand = wb[name]
        hdr = next(cand.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if "Attribute Specification" in [g(hdr, i) for i in range(len(hdr))]:
            ws = cand
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
    group_attrs = OrderedDict()                          # group -> [attrs] (global union)
    spec_def = OrderedDict()                             # spec -> OrderedDict(group -> [attrs])
    mand = {}                                            # (spec,group,attr) -> bool
    path_spec = {}                                       # node_path -> spec name
    spec_families = defaultdict(set)                     # spec -> {paths}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        attr, grp, spec, p = g(r, 5), g(r, 9), g(r, 10), g(r, 4)
        m = g(r, 7).lower() == "yes"
        if not attr or not grp or not spec or spec == "#N/A":
            continue
        group_attrs.setdefault(grp, [])
        if attr not in group_attrs[grp]:
            group_attrs[grp].append(attr)
        sd = spec_def.setdefault(spec, OrderedDict())
        sd.setdefault(grp, [])
        if attr not in sd[grp]:
            sd[grp].append(attr)
        mand[(spec, grp, attr)] = m
        if p:
            path_spec[p] = spec
            spec_families[spec].add(p)
    return group_attrs, spec_def, mand, path_spec, spec_families


def build_attribute_group(spec, spec_def, mand, group_id, attr_by_name):
    out = []
    for grp, attrs in spec_def[spec].items():
        gid = group_id.get(grp)
        if not gid:
            continue
        aids = [{"attributeId": attr_by_name[a.lower()],
                 "isMandatory": 1 if mand.get((spec, grp, a)) else 0}
                for a in attrs if a.lower() in attr_by_name]
        out.append({"attributeGroupId": gid, "attributeIds": aids})
    return out


def main():
    args = sys.argv
    path = args[args.index("--file") + 1] if "--file" in args else None
    workers = int(args[args.index("--workers") + 1]) if "--workers" in args else 16
    dry = "--dry-run" in args
    if not path:
        print("need --file"); sys.exit(1)

    tok = token()
    t0 = time.time()
    group_attrs, spec_def, mand, path_spec, spec_families = parse(path)
    ordered_specs = list(spec_def.keys())
    log(f"parsed: {len(group_attrs)} groups | {len(spec_def)} specs | "
        f"{len(path_spec)} families | first spec={ordered_specs[0]!r}")

    attrs = api("GET", "/product/api/attribute?limit=1000&offset=0&keyword=&count=0", tok=tok)[1].get("data", [])
    attr_by_name = {a["name"].strip().lower(): a["id"] for a in attrs}
    cats = api("GET", "/product/api/category?limit=5000&offset=0&keyword=&status=&name=", tok=tok)[1].get("data", [])
    cat_by_path = {c.get("levels", "").strip(): c["categoryId"] for c in cats}
    existing_groups = {x["name"].strip(): x for x in
                       api("GET", "/product/api/attribute-group?limit=5000&offset=0&keyword=&count=0", tok=tok)[1].get("data", [])}
    existing_specs = sorted(api("GET", "/product/api/specification?limit=5000&offset=0&keyword=&count=0", tok=tok)[1].get("data", []),
                            key=lambda x: x["id"])
    log(f"existing: {len(existing_groups)} groups | {len(existing_specs)} specs | {len(attr_by_name)} attrs | {len(cat_by_path)} cats")
    if not attr_by_name or not existing_specs:
        log("ERROR: could not load admin state (token expired?). Re-login and retry.")
        sys.exit(2)

    if dry:
        log(f"would upsert {len(group_attrs)} groups")
        log(f"would update {min(85, len(existing_specs))} existing specs with first "
            f"{min(85, len(ordered_specs))} new specs")
        log(f"would create {max(0, len(ordered_specs) - len(existing_specs))} specs")
        for i in range(min(3, len(existing_specs))):
            log(f"  update id={existing_specs[i]['id']} {existing_specs[i]['name']!r} -> {ordered_specs[i]!r}")
        return

    # ---- Phase 1: groups upsert ----
    group_id = {n: o["id"] for n, o in existing_groups.items()}
    gsort = [max((o.get("sortOrder") or 0) for o in existing_groups.values())]

    def upsert_group(gn):
        ids = [attr_by_name[a.lower()] for a in group_attrs[gn] if a.lower() in attr_by_name]
        obj = existing_groups.get(gn)
        if obj:
            st, resp = api("PUT", f"/product/api/attribute-group/{obj['id']}",
                           {"name": gn, "sortOrder": obj.get("sortOrder") or 1, "attributeIds": ids}, tok=tok)
            return gn, st, obj["id"]
        with _TL:
            gsort[0] += 1
            so = gsort[0]
        st, resp = api("POST", "/product/api/attribute-group/",
                       {"name": gn, "sortOrder": so, "attributeIds": ids}, tok=tok)
        gid = (resp.get("data") or {}).get("id") if isinstance(resp.get("data"), dict) else None
        return gn, st, gid

    log(f"Phase 1: upsert {len(group_attrs)} groups")
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for gn, st, gid in [f.result() for f in as_completed([ex.submit(upsert_group, gn) for gn in group_attrs])]:
            if gid:
                group_id[gn] = gid
            else:
                log(f"  GROUP FAIL {gn!r}: {st}")

    # ---- Phase 2: update first 85, create rest ----
    spec_id = {}
    n_update = min(85, len(existing_specs), len(ordered_specs))
    updates = [(existing_specs[i]["id"], ordered_specs[i]) for i in range(n_update)]
    log(f"Phase 2a: update {len(updates)} existing specs (first 85 by sheet order)")

    def do_update(item):
        sid, spec = item
        ag = build_attribute_group(spec, spec_def, mand, group_id, attr_by_name)
        st, resp = api("PUT", f"/product/api/specification/{sid}",
                       {"name": spec, "status": 1, "attributeGroup": ag, "attributeGroupId": []}, tok=tok)
        return spec, st, sid

    with ThreadPoolExecutor(max_workers=workers) as ex:
        for spec, st, sid in [f.result() for f in as_completed([ex.submit(do_update, it) for it in updates])]:
            if st in (200, 201):
                spec_id[spec] = sid
            else:
                log(f"  UPDATE FAIL {spec!r}: {st}")

    new_specs = [s for s in ordered_specs[n_update:] if s not in spec_id]
    log(f"Phase 2b: create {len(new_specs)} specs")

    def do_create(spec):
        ag = build_attribute_group(spec, spec_def, mand, group_id, attr_by_name)
        st, resp = api("POST", "/product/api/specification",
                       {"name": spec, "status": 1, "attributeGroup": ag, "attributeGroupId": []}, tok=tok)
        sid = (resp.get("data") or {}).get("id") if isinstance(resp.get("data"), dict) else None
        return spec, st, sid

    done = 0
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for spec, st, sid in [f.result() for f in as_completed([ex.submit(do_create, s) for s in new_specs])]:
            done += 1
            if sid:
                spec_id[spec] = sid
            else:
                log(f"  CREATE FAIL {spec!r}: {st}")
            if done % 250 == 0:
                log(f"  created {done}/{len(new_specs)}")

    # ---- Phase 3: re-assign families ----
    log("Phase 3: re-assign families (set new spec, drop stale specs)")
    tasks = []
    for p, spec in path_spec.items():
        cid = cat_by_path.get(p) or cat_by_path.get(p.replace("Books & General Merchandise", "Toys & General Merchandise"))
        sid = spec_id.get(spec)
        if cid and sid:
            tasks.append((cid, sid))

    def do_assign(task):
        cid, sid = task
        cur = api("GET", f"/product/api/spec-category/{cid}", tok=tok)[1]
        have = [s["id"] for s in (cur.get("data", {}).get("specifications") or [])]
        to_delete = [x for x in have if x != sid]
        if sid in have and not to_delete:
            return cid, True, 0
        st, resp = api("PUT", f"/product/api/spec-category/{cid}",
                       {"specification": [{"specificationId": sid}], "deleteSpecificationIds": to_delete}, tok=tok)
        return cid, st in (200, 201), len(to_delete)

    done = ok = deleted = 0
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for cid, good, ndel in [f.result() for f in as_completed([ex.submit(do_assign, t) for t in tasks])]:
            done += 1
            if good:
                ok += 1
                deleted += ndel
            else:
                log(f"  ASSIGN FAIL cat={cid}")
            if done % 500 == 0:
                log(f"  families {done}/{len(tasks)} (ok {ok}, stale removed {deleted})")

    log(f"assignments ok={ok}/{len(tasks)} | stale specs removed={deleted}")
    log(f"DONE in {(time.time()-t0)/60:.1f} min")


if __name__ == "__main__":
    main()
